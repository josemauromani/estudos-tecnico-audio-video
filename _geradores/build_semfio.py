# -*- coding: utf-8 -*-
"""Gera Microfones-sem-Fio-modelo.xlsx: plano de frequências com alertas, teste de RF, controle de pilhas,
referência e checklist."""
import os, sys
sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
from xlsx_lib import *
from openpyxl.utils import get_column_letter as col_letter

OUT = sys.argv[1]

PURPLE_, PURPLE_T = HUE["purple"]
TEAL, TEAL_T = HUE["teal"]
AMBER, AMBER_T = HUE["amber"]
SIM_NAO = ["Sim", "Não"]
N = 16            # sistemas no plano, no teste e no controle de pilhas
PLANO = "Plano de frequências"

CHECK = [
    ("Pilhas", "Todos os transmissores têm pilhas novas ou baterias com carga completa, e há reservas ao lado da mesa."),
    ("Antenas", "Os receptores estão perto do palco, com as antenas à vista da área de uso e fora de rack fechado."),
    ("Antenas", "As antenas estão em V, sem encostar em metal nem uma na outra, longe de telão de LED, computador e roteador."),
    ("Frequências", "A varredura foi feita com os transmissores desligados e com telão, luz e demais equipamentos ligados."),
    ("Frequências", "Os sistemas da mesma linha usam o mesmo grupo, cada um em um canal diferente."),
    ("Frequências", "Nenhuma frequência se repete, fica perto demais de outra ou cai em canal de TV ativo: a planilha não mostra alerta."),
    ("Frequências", "Cada transmissor foi sincronizado com o seu receptor e identificado com o nome de quem usa."),
    ("Frequências", "Com todos ligados, desligar um transmissor de cada vez deixa o receptor dele sem RF."),
    ("Áudio", "O ganho de cada transmissor foi ajustado com a voz mais forte do evento, sem acender o pico."),
    ("Áudio", "A saída do receptor (mic ou linha) combina com a entrada da mesa, e o ganho do canal foi ajustado."),
    ("Teste", "O teste de caminhada cobriu toda a área de uso, sem falhas."),
    ("Plano B", "Há um microfone com fio ligado, testado e fechado na mesa, e a equipe sabe quem o leva ao palco."),
]

# Exemplo 1 do treinamento: seis sistemas UHF com as frequências escolhidas à mão, de 0,5 em 0,5 MHz
EX_PLANO = {
    "head": {"evento": "Culto de domingo, 19 h", "local": "Igreja de 200 lugares", "data": date(2026, 10, 4),
             "resp": "Operador de som da escala"},
    "par": (0.35, 0.1, 698, 806),
    "tv": [18, 22, 27, 33],
    "sistemas": [  # sistema, usuário, transmissor, faixa, grupo, canal, frequência
        ("Receptor 1 (linha A)", "Pastor", "Headset", "UHF", None, None, 530.000),
        ("Receptor 2 (linha A)", "Cantora", "De mão", "UHF", None, None, 530.500),
        ("Receptor 3 (linha A)", "Voz de apoio 1", "De mão", "UHF", None, None, 531.000),
        ("Receptor 4 (linha A)", "Voz de apoio 2", "De mão", "UHF", None, None, 531.500),
        ("Receptor 5 (linha B)", "Dirigente do louvor", "De mão", "UHF", None, None, 520.300),
        ("Receptor 6 (linha B)", "Reserva e avisos", "De mão", "UHF", None, None, 531.650),
    ],
    "ajustes": [  # sistema, antes, depois, motivo
        ("1. Pastor", "530,000", "534,850", "Era atingido pelo produto dos sistemas 2 e 3 (2 × 530,500 − 531,000 = 530,000)."),
        ("2. Cantora", "530,500", "535,575", "Era atingido pelo produto dos sistemas 3 e 4."),
        ("3. Voz de apoio 1", "531,000", "536,000", "Era atingido pelo produto dos sistemas 2 e 1."),
        ("4. Voz de apoio 2", "531,500", "536,575", "Estava a 0,150 MHz do sistema 6."),
        ("5. Dirigente do louvor", "520,300", "537,475", "Estava dentro do canal de TV 22, ativo na cidade."),
        ("6. Reserva e avisos", "531,650", "537,950", "Estava a 0,150 MHz do sistema 4."),
        ("Resultado", "3 com intermodulação, 2 próximas, 1 em TV", "6 sistemas OK",
         "Novo plano nos canais de TV 24 e 25, livres, com espaçamento desigual. Conferido com todos ligados."),
    ],
}


def hint(ws, ref, text, bg=GRAY, merge=None, h="center"):
    put(ws, ref, text, f=font(9, i=True, c=MUTED), bg=bg, h=h, merge=merge)


def dv_range(ws, rng, origem, prompt):
    """Lista suspensa que lê as opções de um intervalo: aceita opção com vírgula (2,4 GHz)."""
    dv = DataValidation(type="list", formula1=origem, allow_blank=True)
    dv.promptTitle, dv.prompt = "Opções", prompt
    dv.errorTitle, dv.error = "Valor inválido", "Escolha uma das opções da lista."
    dv.showErrorMessage = dv.showInputMessage = True
    ws.add_data_validation(dv)
    dv.add(rng)


def cabecalho(ws, linhas, H, hb, lab=("B", "C"), val=("D", "F")):
    for r, rot, k, fmt in linhas:
        label(ws, f"{lab[0]}{r}", rot, merge=f"{lab[0]}{r}:{lab[1]}{r}")
        inp(ws, f"{val[0]}{r}", H.get(k), merge=f"{val[0]}{r}:{val[1]}{r}", bg=hb, fmt=fmt)
        ws.row_dimensions[r].height = 21.75
        if fmt == DATE:
            dv_date(ws, f"{val[0]}{r}")


# ================================================================== aba Plano de frequências
P_W = {"A": 2, "B": 5, "C": 24, "D": 24, "E": 13, "F": 10, "G": 8, "H": 8, "I": 13, "J": 9, "K": 11, "L": 9, "M": 11,
       "N": 13, "O": 52, "P": 26, "Q": 2}
P1 = 15                 # primeira linha de sistemas
D0, M0 = 18, 35         # primeira coluna da matriz de distâncias (R) e da matriz de produtos (AI)
FAIXAS = ["VHF", "UHF", "2,4 GHz", "Outra"]
SIT_PLANO = [("OK", GREEN), ("Automática", GRAY), ("Repetida", RED), ("Próxima", RED), ("Telefonia", RED),
             ("TV", YELLOW), ("Intermod.", YELLOW)]


def aba_plano(ws, tab, data=None):
    widths(ws, P_W)
    is_ex = data is not None
    title(ws, "Plano de frequências — microfones sem fio",
          "Exemplo preenchido, para consulta. Use a aba Plano de frequências para o seu evento." if is_ex
          else "Preencha as células em amarelo-claro. As células cinza conferem repetição, proximidade, TV e intermodulação.", "P")
    hb = WHITE if is_ex else INPUT
    H = data["head"] if is_ex else {}
    cabecalho(ws, [(4, "Evento", "evento", None), (5, "Local", "local", None), (6, "Data", "data", DATE),
                   (7, "Responsável", "resp", None)], H, hb)
    par = data["par"] if is_ex else (0.35, 0.1, 698, 806)
    for k, (rot, fmt, dica) in enumerate([
            ("Distância mínima entre frequências (MHz)", "0.000", "Ponto de partida para sistemas analógicos. O manual do seu sistema dá o valor certo."),
            ("Margem de intermodulação (MHz)", "0.000", "Distância mínima entre uma frequência e um produto do tipo 2 × A − B."),
            ("Faixa a evitar: de (MHz)", "0", "Faixa da telefonia móvel no Brasil. Confira a regulamentação vigente no seu local."),
            ("Faixa a evitar: até (MHz)", "0", "Deixe as duas células em branco se não quiser esta conferência.")]):
        r = 4 + k
        label(ws, f"G{r}", rot, merge=f"G{r}:K{r}")
        inp(ws, f"L{r}", par[k], h="center", bg=hb, fmt=fmt)
        hint(ws, f"M{r}", dica, merge=f"M{r}:P{r}", h="left")
    dv_number(ws, "L4:L7")
    DMIN, TOL, EV1, EV2 = "$L$4", "$L$5", "$L$6", "$L$7"

    label(ws, "B9", "Canais de TV digital ativos na cidade", merge="B9:D9")
    tv = data["tv"] if is_ex else []
    for k, c in enumerate("EFGHIJKL"):
        inp(ws, f"{c}9", tv[k] if k < len(tv) else None, h="center", bg=hb)
    hint(ws, "M9", "Canal físico (de RF), de 14 a 51: não é o número que aparece na TV. Até oito canais.", merge="M9:P9", h="left")
    dv_number(ws, "E9:L9")
    ws.row_dimensions[9].height = 21.75
    TVS = "$E$9:$L$9"

    band(ws, 11, "Plano de frequências", "P")
    heads = [("B", "#", INK), ("C", "Sistema (receptor)", INK), ("D", "Usuário ou função", INK), ("E", "Transmissor", INK),
             ("F", "Faixa", INK), ("G", "Grupo", PURPLE_), ("H", "Canal", PURPLE_), ("I", "Frequência", PURPLE_),
             ("J", "Canal de TV", TEAL), ("K", "Menor distância", TEAL), ("L", "Do sistema", TEAL), ("M", "Produtos perto", TEAL),
             ("N", "Situação", TEAL), ("O", "O que fazer", TEAL), ("P", "Observação", INK)]
    for c, t, bg in heads:
        head(ws, f"{c}12", t, bg=bg)
    ws.row_dimensions[12].height = 30
    for c, t in zip("BCDEFGHIJKLMNOP", ["nº", "modelo ou etiqueta", "quem usa", "tipo", "", "do visor", "do visor", "MHz",
                                        "6 MHz cada", "MHz", "nº", "intermod.", "", "", ""]):
        hint(ws, f"{c}13", t)
    ws.row_dimensions[13].height = 21.75
    if is_ex:
        hint(ws, "B14", "Frequências escolhidas à mão, de 0,5 em 0,5 MHz, antes da correção. As colunas de J a O são calculadas.",
             merge="B14:P14", h="left")
    else:
        for c, v in zip("BCDEFGHIJKLMNO", ["Ex.", "Receptor 1 (UHF)", "Pastor", "Headset", "UHF", 3, 7, 538.425, 25, 0.75, 2, 0,
                                           "OK", "Manter"]):
            ex(ws, f"{c}14", v, h="left" if c in "CDO" else "center", fmt="0.000" if c in "IK" else None)
        ex(ws, "P14", "")
        note(ws, "I12", "Frequência em MHz, como aparece no visor do receptor: 538,425.\nSistema de 2,4 GHz não tem frequência fixa: deixe em branco.")
        note(ws, "J12", "Canal de TV (de 6 MHz) em que a frequência cai, de 14 a 51.\nCompare com os canais ativos informados no alto.")
        note(ws, "K12", "Distância até a frequência mais próxima de outro sistema da lista.")
        note(ws, "M12", "Quantos produtos de intermodulação (2 × A − B) de outros dois sistemas caem a menos da margem desta frequência.")
        note(ws, "N12", "Repetida: igual à de outro sistema.\nPróxima: abaixo da distância mínima.\nTelefonia: dentro da faixa a evitar.\nTV: dentro de canal de TV ativo.\nIntermod.: há produto perto.")
    ws.row_dimensions[14].height = 24

    P2 = P1 + N - 1
    # listas e cálculos auxiliares, fora da área de impressão
    put(ws, "R3", "Lista de faixas", f=font(9, i=True, c=MUTED), box=False)
    for k, t in enumerate(FAIXAS):
        put(ws, f"R{4 + k}", t, f=font(9, c=MUTED), bg=GRAY)
    put(ws, f"{col_letter(D0)}12", "Cálculo auxiliar (não altere): distância entre os sistemas, em MHz",
        f=font(9, i=True, c=MUTED), box=False)
    put(ws, f"{col_letter(M0)}12", "Cálculo auxiliar (não altere): produtos 2 × A − B, em kHz (A na linha, B na coluna)",
        f=font(9, i=True, c=MUTED), box=False)
    for j in range(N):
        for base in (D0, M0):
            put(ws, f"{col_letter(base + j)}13", j + 1, f=font(9, True, c=MUTED), bg=GRAY, h="center")
            ws.column_dimensions[col_letter(base + j)].width = 9
    ws.column_dimensions[col_letter(D0 + N)].width = 2
    PM = f"${col_letter(M0)}${P1}:${col_letter(M0 + N - 1)}${P2}"

    bg = WHITE if is_ex else INPUT
    for i in range(N):
        r = P1 + i
        s = data["sistemas"][i] if is_ex and i < len(data["sistemas"]) else None
        num(ws, f"B{r}", i + 1)
        inp(ws, f"C{r}", s[0] if s else None, bg=bg)
        inp(ws, f"D{r}", s[1] if s else None, bg=bg)
        for c, j in zip("EFGH", range(2, 6)):
            inp(ws, f"{c}{r}", s[j] if s else None, h="center", bg=bg)
        inp(ws, f"I{r}", s[6] if s else None, h="center", bg=bg, fmt="0.000")
        for j in range(N):
            if j == i:
                put(ws, f"{col_letter(D0 + j)}{r}", None, bg=GRAY)
                put(ws, f"{col_letter(M0 + j)}{r}", None, bg=GRAY)
                continue
            rj = P1 + j
            put(ws, f"{col_letter(D0 + j)}{r}", f'=IF(OR($I{r}="",$I{rj}=""),"",ROUND(ABS($I{r}-$I{rj}),4))',
                f=font(9, c=MUTED), h="center", fmt="0.000")
            put(ws, f"{col_letter(M0 + j)}{r}", f'=IF(OR($I{r}="",$I{rj}=""),"",ROUND((2*$I{r}-$I{rj})*1000,0))',
                f=font(9, c=MUTED), h="center", fmt="0")
        DR = f"${col_letter(D0)}{r}:${col_letter(D0 + N - 1)}{r}"
        PR = f"${col_letter(M0)}{r}:${col_letter(M0 + N - 1)}{r}"
        PC = f"{col_letter(M0 + i)}${P1}:{col_letter(M0 + i)}${P2}"
        lo, hi = f'">"&ROUND((I{r}-{TOL})*1000,0)', f'"<"&ROUND((I{r}+{TOL})*1000,0)'
        conta = lambda rg: f"COUNTIFS({rg},{lo},{rg},{hi})"
        canal_tv = f"INT((I{r}-470)/6)+14"
        calc(ws, f"J{r}", f'=IF(I{r}="","",IF(AND(I{r}>=470,I{r}<698),{canal_tv},"–"))', b=False)
        calc(ws, f"K{r}", f'=IF(OR(I{r}="",COUNT({DR})=0),"",MIN({DR}))', b=False, fmt="0.000")
        calc(ws, f"L{r}", f'=IF(K{r}="","",MATCH(K{r},{DR},0))', b=False)
        calc(ws, f"M{r}", f'=IF(I{r}="","",{conta(PM)}-{conta(PR)}-{conta(PC)})', b=False)
        calc(ws, f"N{r}", f'=IF(I{r}="",IF(AND(C{r}<>"",F{r}="2,4 GHz"),"Automática",""),'
                          f'IF(AND(K{r}<>"",K{r}<0.0005),"Repetida",'
                          f'IF(AND(K{r}<>"",{DMIN}<>"",K{r}<{DMIN}),"Próxima",'
                          f'IF(AND({EV1}<>"",{EV2}<>"",I{r}>={EV1},I{r}<={EV2}),"Telefonia",'
                          f'IF(AND(I{r}>=470,I{r}<698,COUNTIF({TVS},{canal_tv})>0),"TV",'
                          f'IF(M{r}>0,"Intermod.","OK"))))))')
        calc(ws, f"O{r}", f'=IF(N{r}="","",IF(N{r}="Automática","Sem frequência fixa: o sistema escolhe o canal. Afaste o receptor do roteador Wi-Fi",'
                          f'IF(N{r}="Repetida","Mesma frequência do sistema nº "&L{r}&": troque o canal de um dos dois",'
                          f'IF(N{r}="Próxima","Perto demais do sistema nº "&L{r}&": afaste as duas frequências",'
                          f'IF(N{r}="Telefonia","Dentro da faixa a evitar: troque de frequência ou de equipamento",'
                          f'IF(N{r}="TV","Dentro do canal de TV "&J{r}&", ativo na cidade: use um canal de TV livre",'
                          f'IF(N{r}="Intermod.","Há "&M{r}&" produto(s) de intermodulação perto: troque esta frequência ou use o grupo do fabricante",'
                          f'"Manter")))))))', h="left", b=False, sz=9)
        inp(ws, f"P{r}", bg=bg)
        ws.row_dimensions[r].height = 21.75
    dv_list(ws, f"E{P1}:E{P2}", ["De mão", "Lapela", "Headset", "Instrumento", "In-ear (IEM)", "Outro"],
            "De mão, Lapela, Headset, Instrumento, In-ear (IEM) ou Outro.")
    dv_range(ws, f"F{P1}:F{P2}", "$R$4:$R$7", "VHF, UHF, 2,4 GHz ou Outra.")
    dv_number(ws, f"I{P1}:I{P2}")
    cf_equal(ws, f"N{P1}:N{P2}", SIT_PLANO)

    s = P2 + 2
    band(ws, s, "Resumo automático", "P")
    cont = lambda t: f'COUNTIF(N{P1}:N{P2},"{t}")'
    linhas = [
        ("Sistemas listados", f"=COUNTA(C{P1}:C{P2})"),
        ("Com frequência informada", f"=COUNT(I{P1}:I{P2})"),
        ("Sem alerta (OK)", f'={cont("OK")}'),
        ("Repetidas ou próximas demais", f'={cont("Repetida")}+{cont("Próxima")}'),
        ("Em canal de TV ativo ou na faixa a evitar", f'={cont("TV")}+{cont("Telefonia")}'),
        ("Com risco de intermodulação", f'={cont("Intermod.")}'),
        ("Em 2,4 GHz, com frequência automática", f'={cont("Automática")}'),
    ]
    for k, (t, f) in enumerate(linhas, 1):
        summary(ws, s + k, t, f, "H", "I")
    a = s + len(linhas) + 1
    label(ws, f"B{a}", "Aviso", merge=f"B{a}:H{a}", h="right")
    calc(ws, f"I{a}", f'=IF(I{s + 1}+I{s + 2}=0,"Liste os sistemas e as frequências",'
                      f'IF(I{s + 4}>0,"Há frequências repetidas ou próximas demais: corrija antes de ligar os transmissores",'
                      f'IF(I{s + 5}>0,"Há frequência em canal de TV ativo ou na faixa a evitar: troque",'
                      f'IF(I{s + 6}>0,"Há risco de intermodulação: troque as frequências marcadas ou use o grupo do fabricante","OK"))))',
         b=False, merge=f"I{a}:O{a}", h="left")
    ws.row_dimensions[a].height = 21.75
    cf_ok(ws, f"I{a}:O{a}")
    hint(ws, f"B{a + 1}", "A conferência cobre só os produtos de terceira ordem de dois transmissores (2 × A − B). Ela não substitui os grupos "
                          "do fabricante, o programa de coordenação nem o teste com todos os transmissores ligados.",
         merge=f"B{a + 1}:P{a + 1}", h="left")
    ws.row_dimensions[a + 1].height = 27
    end = a + 1

    if is_ex:
        band(ws, end + 2, "O que foi ajustado", "P")
        r = end + 3
        put(ws, f"B{r}", "Sistema", f=font(10, True), bg=GRAY, h="center", merge=f"B{r}:C{r}")
        put(ws, f"D{r}", "Antes (MHz)", f=font(10, True), bg=GRAY, h="center", merge=f"D{r}:F{r}")
        put(ws, f"G{r}", "Depois (MHz)", f=font(10, True), bg=GRAY, h="center", merge=f"G{r}:I{r}")
        put(ws, f"J{r}", "Motivo", f=font(10, True), bg=GRAY, h="center", merge=f"J{r}:P{r}")
        ws.row_dimensions[r].height = 21.75
        for sis, antes, depois, motivo in data["ajustes"]:
            r += 1
            put(ws, f"B{r}", sis, f=font(10, True), bg=GRAY, merge=f"B{r}:C{r}")
            put(ws, f"D{r}", antes, h="center", merge=f"D{r}:F{r}")
            put(ws, f"G{r}", depois, h="center", merge=f"G{r}:I{r}")
            put(ws, f"J{r}", motivo, merge=f"J{r}:P{r}")
            ws.row_dimensions[r].height = 21.75
        end = r
    ws.freeze_panes = f"E{P1}"
    setup(ws, tab, f"B1:P{end}", fit_height=True)


# ================================================================== aba Teste de RF
T_W = {"A": 2, "B": 5, "C": 24, "D": 22, "E": 12, "F": 16, "G": 18, "H": 14, "I": 14, "J": 14, "K": 14, "L": 54, "M": 28, "N": 2}
T1 = 11


def aba_teste(ws, tab):
    widths(ws, T_W)
    title(ws, "Teste de RF — sistema por sistema",
          "Os sistemas vêm da aba Plano de frequências. Responda Sim ou Não a cada teste, com o sistema montado no local.", "M")
    cabecalho(ws, [(4, "Evento", "evento", None), (5, "Data", "data", DATE), (6, "Quem testou", "resp", None)], {}, INPUT,
              val=("D", "E"))
    hint(ws, "F4", "Teste 1: todos os transmissores desligados, receptores ligados, telão e luz ligados.", merge="F4:M4", h="left")
    hint(ws, "F5", "Teste 2: todos os transmissores ligados e afastados uns dos outros; desligue um de cada vez e olhe o receptor dele.",
         merge="F5:M5", h="left")
    hint(ws, "F6", "Testes 3 a 5: uma pessoa anda por toda a área de uso falando forte; outra olha o receptor e ouve no fone.",
         merge="F6:M6", h="left")

    band(ws, 8, "Registro dos testes", "M")
    for c, t, bg in [("B", "#", INK), ("C", "Sistema", INK), ("D", "Usuário", INK), ("E", "Frequência", INK),
                     ("F", "1. Transmissores desligados: mostra RF?", PURPLE_),
                     ("G", "2. Só este desligado, os outros ligados: mostra RF?", PURPLE_),
                     ("H", "3. Caminhada: houve falha?", PURPLE_), ("I", "4. Antenas A e B recebem?", PURPLE_),
                     ("J", "5. Pico de áudio acende?", PURPLE_),
                     ("K", "Situação", TEAL), ("L", "O que fazer", TEAL), ("M", "Observação", INK)]:
        head(ws, f"{c}9", t, bg=bg)
    ws.row_dimensions[9].height = 57
    for c, t in zip("BCDEFGHIJKLM", ["", "do plano", "do plano", "MHz", "esperado: Não", "esperado: Não", "esperado: Não",
                                     "esperado: Sim", "esperado: Não", "", "", ""]):
        hint(ws, f"{c}10", t)
    ws.row_dimensions[10].height = 21.75
    note(ws, "F9", "Com o transmissor deste sistema desligado, o medidor de RF do receptor deve ficar vazio.\nSe mostra sinal, há outra emissão na frequência.")
    note(ws, "G9", "Com os outros transmissores ligados e só este desligado, o receptor deve continuar sem RF.\nSe mostra sinal, é intermodulação ou frequência próxima.")
    note(ws, "I9", "Os indicadores A e B devem alternar ou acender durante a caminhada.\nSe só um acende, confira a antena e o cabo do outro.")
    T2 = T1 + N - 1
    for i in range(N):
        r, pr = T1 + i, P1 + i
        num(ws, f"B{r}", i + 1)
        for c, pc in (("C", "C"), ("D", "D"), ("E", "I")):
            calc(ws, f"{c}{r}", f"=IF('{PLANO}'!{pc}{pr}=\"\",\"\",'{PLANO}'!{pc}{pr})", h="left" if c != "E" else "center",
                 b=False, fmt="0.000" if c == "E" else None)
        for c in "FGHIJ":
            inp(ws, f"{c}{r}", h="center")
        calc(ws, f"K{r}", f'=IF(COUNTA(F{r}:J{r})=0,"",IF(F{r}="Sim","Interferência",IF(G{r}="Sim","Intermod.",'
                          f'IF(I{r}="Não","Antena",IF(H{r}="Sim","Falha",IF(J{r}="Sim","Ganho alto",'
                          f'IF(COUNTA(F{r}:J{r})<5,"Incompleto","OK")))))))')
        calc(ws, f"L{r}", f'=IF(K{r}="","",IF(K{r}="Interferência","Outra emissão na frequência: refaça a varredura e troque o canal",'
                          f'IF(K{r}="Intermod.","Produto de outros transmissores: refaça o plano com o grupo do fabricante",'
                          f'IF(K{r}="Antena","Uma antena não recebe: confira conector, cabo e posição",'
                          f'IF(K{r}="Falha","Dê visada às antenas, aproxime o receptor do palco e repita a caminhada",'
                          f'IF(K{r}="Ganho alto","Baixe o ganho do transmissor até o pico apagar na voz mais forte",'
                          f'IF(K{r}="Incompleto","Faça os cinco testes","Manter")))))))', h="left", b=False, sz=9)
        inp(ws, f"M{r}")
        ws.row_dimensions[r].height = 21.75
    dv_list(ws, f"F{T1}:J{T2}", SIM_NAO, "Sim ou Não.")
    cf_status(ws, f"K{T1}:K{T2}", f"K{T1}")

    s = T2 + 2
    band(ws, s, "Resumo automático", "M")
    prob = "+".join(f'COUNTIF(K{T1}:K{T2},"{t}")' for t in ("Interferência", "Intermod.", "Antena", "Falha", "Ganho alto"))
    summary(ws, s + 1, "Sistemas no plano", f"=COUNTA('{PLANO}'!C{P1}:C{P1 + N - 1})", "E", "F")
    summary(ws, s + 2, "Testados e OK", f'=COUNTIF(K{T1}:K{T2},"OK")', "E", "F")
    summary(ws, s + 3, "Com problema", f"={prob}", "E", "F")
    summary(ws, s + 4, "Com teste incompleto", f'=COUNTIF(K{T1}:K{T2},"Incompleto")', "E", "F")
    a = s + 5
    label(ws, f"B{a}", "Aviso", merge=f"B{a}:E{a}", h="right")
    calc(ws, f"F{a}", f'=IF(F{s + 1}=0,"Preencha primeiro a aba Plano de frequências",'
                      f'IF(F{s + 3}>0,"Há sistema com problema: corrija e repita o teste",'
                      f'IF(F{s + 2}<F{s + 1},"Há sistema sem teste completo","OK")))',
         b=False, merge=f"F{a}:L{a}", h="left")
    ws.row_dimensions[a].height = 21.75
    cf_ok(ws, f"F{a}:L{a}")
    ws.freeze_panes = f"F{T1}"
    setup(ws, tab, f"B1:M{a}", fit_height=True)


# ================================================================== aba Pilhas e baterias
B_W = {"A": 2, "B": 5, "C": 28, "D": 20, "E": 8, "F": 11, "G": 12, "H": 12, "I": 12, "J": 12, "K": 10, "L": 13, "M": 50,
       "N": 14, "O": 26, "P": 2}
B1 = 12
SIT_PILHA = [("OK", GREEN), ("Atenção", YELLOW), ("Sem dados", YELLOW), ("Trocar", RED)]


def aba_pilhas(ws, tab):
    widths(ws, B_W)
    title(ws, "Pilhas e baterias — controle por transmissor",
          "Anote a autonomia e o tempo de uso de cada transmissor. A planilha diz se a carga chega ao fim do evento com margem.", "O")
    cabecalho(ws, [(4, "Evento", "evento", None), (5, "Data", "data", DATE), (6, "Responsável", "resp", None)], {}, INPUT,
              val=("D", "E"))
    for r, rot, v, fmt, dica in [
            (4, "Horas ligado no evento (passagem + evento)", 4, "0.0", "Conte desde a hora em que o transmissor é ligado para a passagem de som."),
            (5, "Fator de margem", 1.5, "0.0", "A carga restante deve cobrir as horas do evento vezes este fator. Regra prática deste material.")]:
        label(ws, f"F{r}", rot, merge=f"F{r}:J{r}")
        inp(ws, f"K{r}", v, h="center", fmt=fmt)
        hint(ws, f"L{r}", dica, merge=f"L{r}:O{r}", h="left")
    dv_number(ws, "K4:K5")
    hint(ws, "F6", "Pilha recarregável antiga dura menos que o manual diz: meça a autonomia real em um ensaio e use esse valor.",
         merge="F6:O6", h="left")
    HORAS, FATOR = "$K$4", "$K$5"

    band(ws, 8, "Controle de pilhas e baterias", "O")
    for c, t, bg in [("B", "#", INK), ("C", "Transmissor ou usuário", INK), ("D", "Tipo de pilha", INK), ("E", "Qtde", INK),
                     ("F", "Autonomia", AMBER), ("G", "Nova ou carga completa?", AMBER), ("H", "Horas já usadas", AMBER),
                     ("I", "Horas neste evento", AMBER), ("J", "Carga restante", TEAL), ("K", "Folga", TEAL),
                     ("L", "Situação", TEAL), ("M", "O que fazer", TEAL), ("N", "Última troca ou carga", INK), ("O", "Observação", INK)]:
        head(ws, f"{c}9", t, bg=bg)
    ws.row_dimensions[9].height = 30
    for c, t in zip("BCDEFGHIJKLMNO", ["nº", "", "", "por aparelho", "horas", "Sim / Não", "desde a troca", "vazio: valor do alto",
                                       "horas", "horas", "", "", "data", ""]):
        hint(ws, f"{c}10", t)
    ws.row_dimensions[10].height = 27
    for c, v in zip("BCDEFGHIJKLM", ["Ex.", "Headset do pastor", "Alcalina", 2, 8, "Não", 3, "", 5, 1, "Atenção",
                                     "Chega ao fim, mas sem a margem: troque ou deixe a reserva ao lado"]):
        ex(ws, f"{c}11", v, h="left" if c in "CDM" else "center")
    ex(ws, "N11", "")
    ex(ws, "O11", "")
    ws.row_dimensions[11].height = 24
    note(ws, "F9", "Horas de funcionamento com pilha nova ou carga completa.\nUse o valor do manual do transmissor ou, melhor, o que você mediu.")
    note(ws, "G9", "Sim: a pilha foi trocada ou a bateria saiu agora do carregador. As horas já usadas contam como zero.")
    note(ws, "H9", "Horas de uso desde a última troca ou carga: ensaios, passagens e eventos anteriores.")
    note(ws, "L9", "Trocar: a carga não chega ao fim do evento.\nAtenção: chega, mas sem a margem.\nSem dados: falta a autonomia ou o tempo de uso.")
    B2 = B1 + N - 1
    for i in range(N):
        r = B1 + i
        num(ws, f"B{r}", i + 1)
        inp(ws, f"C{r}")
        for c in "DEFGHI":
            inp(ws, f"{c}{r}", h="center")
        need = f'IF(I{r}="",{HORAS},I{r})'
        calc(ws, f"J{r}", f'=IF(OR(F{r}="",AND(G{r}<>"Sim",H{r}="")),"",MAX(0,F{r}-IF(G{r}="Sim",0,H{r})))', b=False, fmt="0.0")
        calc(ws, f"K{r}", f'=IF(OR(J{r}="",{need}=""),"",J{r}-{need})', b=False, fmt="+0.0;-0.0;0.0")
        calc(ws, f"L{r}", f'=IF(AND(C{r}="",F{r}=""),"",IF(OR(J{r}="",K{r}=""),"Sem dados",IF(K{r}<0,"Trocar",'
                          f'IF(J{r}<{need}*{FATOR},"Atenção","OK"))))')
        calc(ws, f"M{r}", f'=IF(L{r}="","",IF(L{r}="Sem dados","Informe a autonomia e as horas já usadas",'
                          f'IF(L{r}="Trocar","Não chega ao fim do evento: troque ou recarregue agora",'
                          f'IF(L{r}="Atenção","Chega ao fim, mas sem a margem: troque ou deixe a reserva ao lado","Manter"))))',
             h="left", b=False, sz=9)
        inp(ws, f"N{r}", h="center", fmt=DATE)
        inp(ws, f"O{r}")
        ws.row_dimensions[r].height = 21.75
    dv_list(ws, f"D{B1}:D{B2}", ["Alcalina", "NiMH recarregável", "Lítio recarregável", "Lítio descartável", "Outra"],
            "Alcalina, NiMH recarregável, Lítio recarregável, Lítio descartável ou Outra.")
    dv_list(ws, f"G{B1}:G{B2}", SIM_NAO, "A pilha é nova, ou a bateria saiu agora do carregador?")
    dv_number(ws, f"E{B1}:F{B2}")
    dv_number(ws, f"H{B1}:I{B2}")
    dv_date(ws, f"N{B1}:N{B2}")
    cf_equal(ws, f"L{B1}:L{B2}", SIT_PILHA)

    s = B2 + 2
    band(ws, s, "Resumo automático", "O")
    cont = lambda t: f'COUNTIF(L{B1}:L{B2},"{t}")'
    summary(ws, s + 1, "Transmissores listados", f"=COUNTA(C{B1}:C{B2})", "E", "F")
    summary(ws, s + 2, "Com carga e margem (OK)", f'={cont("OK")}', "E", "F")
    summary(ws, s + 3, "Sem a margem (Atenção)", f'={cont("Atenção")}', "E", "F")
    summary(ws, s + 4, "A trocar ou recarregar", f'={cont("Trocar")}', "E", "F")
    summary(ws, s + 5, "Pilhas a separar para troca e reserva",
            f'=SUMIF(L{B1}:L{B2},"Trocar",E{B1}:E{B2})+SUMIF(L{B1}:L{B2},"Atenção",E{B1}:E{B2})', "E", "F")
    a = s + 6
    label(ws, f"B{a}", "Aviso", merge=f"B{a}:E{a}", h="right")
    calc(ws, f"F{a}", f'=IF(F{s + 1}=0,"Liste os transmissores",IF(F{s + 4}>0,"Há transmissor que não chega ao fim do evento: troque agora",'
                      f'IF({cont("Sem dados")}>0,"Há transmissor sem dados de carga",'
                      f'IF(F{s + 3}>0,"Há transmissor sem a margem: tenha reservas à mão","OK"))))',
         b=False, merge=f"F{a}:M{a}", h="left")
    ws.row_dimensions[a].height = 21.75
    cf_ok(ws, f"F{a}:M{a}")
    ws.freeze_panes = f"D{B1}"
    setup(ws, tab, f"B1:O{a}", fit_height=True)


# ================================================================== pasta de trabalho
wb = Workbook()
wb.properties.title = "Microfones sem fio — plano de frequências e controle de pilhas"
wb.properties.creator = "Estudos: Técnico em áudio e vídeo"
wb.properties.language = "pt-BR"

instrucoes(wb, "Microfones sem fio — como usar esta planilha",
           "Plano de frequências com alertas, teste de RF por sistema e controle de pilhas, para usar na montagem do evento.", [
    ("Legenda: onde preencher", [
        ("Amarelo-claro", "Células de entrada. É aqui que você digita."),
        ("Cinza", "Células calculadas ou fixas (rótulos, contas, avisos). Não altere."),
        ("Cinza em itálico", "Linha de exemplo (marcada com “Ex.”) e dicas de preenchimento. Não entram nas contas."),
        ("Listas suspensas", "Transmissor, faixa, tipo de pilha e as respostas Sim ou Não aceitam só as opções da lista."),
        ("Frequências", "Digite a frequência em MHz, como aparece no visor do receptor: 538,425. Não escreva a unidade."),
        ("Comentários", "Células com triângulo vermelho no canto trazem uma dica. Passe o mouse sobre elas."),
    ]),
    ("Ordem de preenchimento", [
        ("1", "Aba Plano de frequências: preencha o cabeçalho e os canais de TV digital ativos na sua cidade (canal físico)."),
        ("2", "No local: com os transmissores desligados, faça a varredura (scan) em cada receptor e escolha o grupo e os canais."),
        ("3", "Aba Plano de frequências: liste cada sistema, quem usa, o grupo, o canal e a frequência do visor."),
        ("4", "Veja a situação de cada linha. Troque as frequências com alerta e digite as novas, até o aviso ficar em OK."),
        ("5", "Aba Pilhas e baterias: anote a autonomia e o tempo de uso de cada transmissor. Troque o que a planilha mandar trocar."),
        ("6", "Aba Teste de RF: com tudo montado, faça os cinco testes de cada sistema e responda Sim ou Não."),
        ("7", "Aba Checklist: confira os doze itens antes de abrir as portas."),
    ]),
    ("Abas da planilha", [
        ("Plano de frequências", "Até 16 sistemas, com alerta de frequência repetida, próxima demais, em canal de TV ativo, na faixa a evitar ou com intermodulação."),
        ("Teste de RF", "Os cinco testes de cada sistema do plano, com a situação e o que fazer em cada caso."),
        ("Pilhas e baterias", "Carga restante e folga de cada transmissor para as horas do evento, com a quantidade de pilhas a separar."),
        ("Referência", "Faixas, valores de partida da coordenação, antenas, pilhas e tabela de sintomas."),
        ("Checklist", "Os doze itens do treinamento, com percentual atendido."),
        ("Exemplo - Plano", "Seis sistemas com as frequências escolhidas à mão, antes da correção (exemplo 1 do treinamento)."),
    ]),
    ("Regras e premissas do modelo", [
        ("Valores de partida", "Distância mínima de 0,35 MHz entre frequências e margem de 0,1 MHz para intermodulação são ponto de partida para sistemas analógicos. Confira o manual do seu sistema e troque os valores no alto da aba."),
        ("Canal de TV", "A conta supõe canais de 6 MHz, do 14 (470 MHz) ao 51 (698 MHz), como na TV digital brasileira. Informe o canal físico, que não é o número mostrado na TV."),
        ("Faixa a evitar", "Vem preenchida com a faixa da telefonia móvel no Brasil (698 a 806 MHz). As regras mudam: confira a regulamentação vigente no seu local."),
        ("Intermodulação", "A planilha calcula os produtos de terceira ordem de dois transmissores (2 × A − B) e conta os que caem perto de cada frequência. Produtos de três transmissores ficam de fora."),
        ("Limite da conferência", "O plano sem alerta é um bom começo, e não uma garantia. Os grupos do fabricante, o programa de coordenação e o teste com todos os transmissores ligados continuam valendo."),
        ("Sistemas de 2,4 GHz", "Não têm frequência fixa: deixe a frequência em branco e escolha a faixa 2,4 GHz. A situação aparece como Automática."),
        ("Pilhas", "Situação OK quando a carga restante cobre as horas do evento vezes o fator de margem. O fator de 1,5 é regra prática deste material."),
        ("Exemplo", "Os dados da aba de exemplo são ilustrativos, criados para este material."),
    ]),
])

aba_plano(wb.create_sheet(PLANO), PURPLE_)
aba_teste(wb.create_sheet("Teste de RF"), TEAL)
aba_pilhas(wb.create_sheet("Pilhas e baterias"), AMBER)

# ------------------------------------------------------------------ Referência
ws = wb.create_sheet("Referência")
widths(ws, {"A": 2, "B": 34, "C": 30, "D": 34, "E": 62, "F": 2})
title(ws, "Referência — valores típicos", "Valores de referência, para servir de ponto de partida. Confira o manual do seu equipamento.", "E")


def tabela(r, titulo, cab, linhas, alt=31.5):
    band(ws, r, titulo, "E")
    r += 1
    for c, t in zip("BCDE", cab):
        put(ws, f"{c}{r}", t, f=font(10, True), bg=GRAY, h="center")
    ws.row_dimensions[r].height = 21.75
    for lin in linhas:
        r += 1
        put(ws, f"B{r}", lin[0], f=font(10, True))
        for c, t in zip("CDE", lin[1:]):
            put(ws, f"{c}{r}", t, h="left" if c == "E" else "center")
        ws.row_dimensions[r].height = alt
    return r + 2


r = tabela(4, "Faixas de operação", ["Faixa", "Frequências típicas", "Ponto forte", "Cuidado"], [
    ("VHF", "174 a 216 MHz", "Custo baixo; a onda longa contorna melhor os obstáculos", "Poucas frequências, muitas vezes fixas. Mais sujeita a ruído elétrico. Poucos sistemas juntos."),
    ("UHF", "470 a 698 MHz", "Muitas frequências; é a faixa dos sistemas profissionais", "Divide o espaço com a TV digital: os canais livres mudam de cidade para cidade."),
    ("Faixa de 700 MHz", "698 a 806 MHz", "–", "Usada pela telefonia móvel no Brasil. Não use equipamento antigo dessa faixa."),
    ("2,4 GHz", "2.400 a 2.483,5 MHz", "Escolhe o canal sozinho; uso liberado em quase todo o mundo", "Divide o espaço com Wi-Fi e Bluetooth. Alcance menor. Em geral, de 4 a 8 sistemas juntos."),
])
r = tabela(r, "Coordenação: valores de partida", ["Regra", "Valor de partida", "Depende de", "Observação"], [
    ("Distância entre duas frequências", "0,3 a 0,4 MHz ou mais", "modelo; digitais aceitam menos", "Nunca duas frequências iguais. O manual do sistema dá o valor mínimo."),
    ("Distância de um produto 2 × A − B", "0,1 MHz ou mais", "qualidade do receptor", "Espaçamento igual entre três frequências coloca o produto em cima da terceira."),
    ("Canal de TV digital ativo", "fora dos 6 MHz inteiros", "cidade", "O canal 14 vai de 470 a 476 MHz; cada canal seguinte fica 6 MHz acima."),
    ("Sistemas da mesma linha", "mesmo grupo, canais diferentes", "fabricante", "Grupos de linhas ou de marcas diferentes não combinam entre si."),
    ("Transmissores de in-ear", "bloco próprio, a alguns MHz", "fabricante", "Entram no mesmo plano. As antenas deles ficam longe das antenas dos receptores."),
])
r = tabela(r, "Antenas e posição", ["Item", "Valor de partida", "Por quê", "Observação"], [
    ("Visada", "antena à vista do transmissor", "o corpo e o metal barram o rádio", "Antenas acima da cabeça do público: 2 m de altura ou mais."),
    ("Transmissor até a antena", "de 3 m a cerca de 30 m", "perto demais satura; longe demais falha", "O alcance do manual vale em campo aberto. Em sala cheia, conte com bem menos."),
    ("Antenas do receptor", "em V, abertas a 45° cada", "diversidade", "Sem encostar uma na outra nem em metal. Nunca dentro de rack fechado."),
    ("Antenas remotas", "afastadas pelo menos 30 a 50 cm", "diversidade", "Cabo coaxial de 50 ohms, o mais curto possível."),
    ("Cabo de antena fino (tipo RG-58)", "perde de 0,3 a 0,5 dB por metro", "frequência e cabo", "Em UHF. Dez metros perdem de 3 a 5 dB. Cabo de baixa perda: 0,1 a 0,2 dB por metro. Prefira esticar o cabo de áudio."),
    ("Fontes de ruído", "a 1 ou 2 m, no mínimo", "ruído de rádio em banda larga", "Telão de LED, processador de vídeo, computador, roteador, mesa digital, fontes chaveadas."),
])
r = tabela(r, "Pilhas e baterias", ["Tipo", "Autonomia típica", "Medidor do transmissor", "Observação"], [
    ("Alcalina AA", "6 a 14 h, conforme o transmissor", "confiável", "Nova a cada evento importante. As usadas ficam para os ensaios."),
    ("NiMH recarregável AA", "parecida quando nova; cai com a idade", "marca menos e cai de repente", "Escolha o tipo de pilha no menu do transmissor, se houver. Par da mesma idade."),
    ("Bateria de lítio do fabricante", "8 h ou mais", "preciso, muitas vezes em horas", "Use só o carregador do sistema."),
    ("9 V (transmissores antigos)", "4 a 8 h com alcalina", "confiável com alcalina", "A recarregável de 9 V dura bem menos."),
])
r = tabela(r, "Sintoma e causa provável", ["Sintoma", "Causa provável", "Primeiro teste", "O que fazer"], [
    ("Queda rápida quando a pessoa anda", "ponto surdo ou corpo na frente", "olhar o RF durante a caminhada", "Dar visada às antenas e aproximar o receptor do palco."),
    ("Ruído com o transmissor desligado", "outra emissão na frequência", "desligar o transmissor e olhar o RF", "Nova varredura e troca de canal. Conferir squelch e tom piloto."),
    ("Falha quando outro microfone liga", "frequências próximas ou intermodulação", "desligar os transmissores um a um", "Refazer o plano com o grupo do fabricante."),
    ("Alcance curto em todo lugar", "canal de TV, antena ou cabo", "trocar de frequência e conferir as antenas", "Sair do canal de TV; tirar as antenas do rack."),
    ("Distorção nas partes fortes", "ganho do transmissor alto", "olhar o pico de áudio do receptor", "Baixar o ganho do transmissor. Depois, refazer o ganho da mesa."),
    ("Chiado de fundo constante", "ganho do transmissor baixo ou RF fraco", "olhar os medidores de AF e de RF", "Subir o ganho do transmissor; conferir a saída mic ou linha."),
    ("Some depois de uma ou duas horas", "pilha", "olhar o indicador e as horas de uso", "Trocar a pilha. Controlar as horas na planilha."),
], alt=27)
hint(ws, f"B{r - 1}", "Faixas e valores típicos, que variam com o fabricante, o modelo e o local. As regras de uso do espectro mudam: "
                      "confira a regulamentação vigente e o manual do seu equipamento.", merge=f"B{r - 1}:E{r - 1}", h="left")
ws.row_dimensions[r - 1].height = 27
setup(ws, HUE["slate"][0], f"B1:E{r - 1}", fit_height=True)

# ------------------------------------------------------------------ Checklist e exemplo
checklist(wb.create_sheet("Checklist"), "Checklist de RF",
          "Confira na montagem, antes de abrir as portas. Escolha a situação de cada item na lista suspensa.",
          CHECK, tab=HUE["blue"][0])
aba_plano(wb.create_sheet("Exemplo - Plano"), MUTED, EX_PLANO)

wb.active = 1
wb.save(OUT)
print("ok", OUT)
