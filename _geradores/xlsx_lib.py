# -*- coding: utf-8 -*-
"""Funções e estilos compartilhados pelos geradores de planilha (build_<nome>.py).

Uso:  from xlsx_lib import *
"""
import sys
from datetime import date

from openpyxl import Workbook
from openpyxl.chart import LineChart, Reference
from openpyxl.chart.layout import Layout, ManualLayout
from openpyxl.chart.shapes import GraphicalProperties
from openpyxl.comments import Comment
from openpyxl.drawing.line import LineProperties
from openpyxl.formatting.rule import CellIsRule, FormulaRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.worksheet.pagebreak import Break

INK, MUTED, GRAY, INPUT, LINE, WHITE = "16242E", "5B6B76", "EEF1F3", "FFF9DB", "C9D2D8", "FFFFFF"
GREEN, YELLOW, RED = "D8EFDD", "FBEBC8", "F6D9D7"
# cor forte e cor clara de cada matiz, iguais às do treinamento
HUE = {
    "blue": ("2B5C8A", "DCE8F3"),
    "amber": ("A96A12", "F6E8CF"),
    "teal": ("1E7B73", "D9EEEB"),
    "purple": ("7A4A9A", "E9DEF1"),
    "red": ("B0413E", "F5DEDC"),
    "slate": ("5B6B76", "E3E8EB"),
}
PURPLE = "7A4A9A"
DATE = "dd/mm/yyyy"

side = Side(style="thin", color=LINE)
BOX = Border(left=side, right=side, top=side, bottom=side)


def fill(c):
    return PatternFill("solid", fgColor=c) if c else PatternFill(fill_type=None)


def font(sz=10, b=False, i=False, c=INK):
    return Font(name="Arial", size=sz, bold=b, italic=i, color=c)


def put(ws, ref, value=None, *, f=None, bg=None, h="left", box=True, fmt=None, merge=None):
    """Escreve e formata uma célula; com merge, formata todo o intervalo."""
    cell = ws[ref]
    if value is not None:
        cell.value = value
    rng = ws[merge] if merge else ((cell,),)
    for row in rng:
        for c in row:
            c.font = f or font()
            c.fill = fill(bg)
            c.alignment = Alignment(horizontal=h, vertical="center", wrap_text=True, indent=1 if h == "left" else 0)
            if box:
                c.border = BOX
            if fmt:
                c.number_format = fmt
    if merge:
        ws.merge_cells(merge)
    return cell


def title(ws, text, sub, last_col):
    put(ws, "B1", text, f=font(16, True, c=WHITE), bg=INK, box=False, merge=f"B1:{last_col}1")
    put(ws, "B2", sub, f=font(10, i=True, c=MUTED), box=False, merge=f"B2:{last_col}2")
    ws.row_dimensions[1].height = 31.5
    ws.row_dimensions[2].height = 19.5


def band(ws, row, text, last_col, color=INK, sz=10):
    put(ws, f"B{row}", text, f=font(sz, True, c=WHITE), bg=color, box=False, merge=f"B{row}:{last_col}{row}")
    ws.row_dimensions[row].height = 21.75


def label(ws, ref, text, merge=None, h="left"):
    put(ws, ref, text, f=font(10, True), bg=GRAY, h=h, merge=merge)


def head(ws, ref, text, bg=INK, fg=WHITE):
    put(ws, ref, text, f=font(10, True, c=fg), bg=bg, h="center")


def num(ws, ref, n):
    put(ws, ref, n, f=font(10, True, c=MUTED), bg=GRAY, h="center")


def ex(ws, ref, text, h="left", fmt=None):
    put(ws, ref, text, f=font(9, i=True, c=MUTED), bg=GRAY, h=h, fmt=fmt)


def calc(ws, ref, formula, h="center", b=True, fmt=None, merge=None, sz=10):
    put(ws, ref, formula, f=font(sz, b), bg=GRAY, h=h, fmt=fmt, merge=merge)


def inp(ws, ref, value=None, h="left", fmt=None, merge=None, bg=INPUT):
    put(ws, ref, value, bg=bg, h=h, fmt=fmt, merge=merge)


def widths(ws, spec):
    for col, w in spec.items():
        ws.column_dimensions[col].width = w


def setup(ws, tab, area, landscape=True, fit_height=False):
    ws.sheet_view.showGridLines = False
    ws.sheet_properties.tabColor = tab
    ws.page_setup.orientation = "landscape" if landscape else "portrait"
    ws.page_setup.paperSize = 9
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 1 if fit_height else 0
    ws.page_margins.left = ws.page_margins.right = 0.4
    ws.page_margins.top = ws.page_margins.bottom = 0.5
    ws.print_options.horizontalCentered = True
    ws.print_area = area


def dv_list(ws, rng, options, prompt):
    dv = DataValidation(type="list", formula1='"%s"' % ",".join(options), allow_blank=True)
    dv.promptTitle, dv.prompt = "Opções", prompt
    dv.errorTitle, dv.error = "Valor inválido", "Escolha uma das opções da lista."
    dv.showErrorMessage = dv.showInputMessage = True
    ws.add_data_validation(dv)
    dv.add(rng)


def dv_date(ws, rng):
    dv = DataValidation(type="date", operator="greaterThanOrEqual", formula1="36526", allow_blank=True)
    dv.errorTitle, dv.error = "Data inválida", "Digite uma data no formato dia/mês/ano."
    dv.showErrorMessage = True
    ws.add_data_validation(dv)
    dv.add(rng)


def dv_number(ws, rng):
    dv = DataValidation(type="decimal", operator="between", formula1="-1000000000", formula2="1000000000", allow_blank=True)
    dv.errorTitle, dv.error = "Valor inválido", "Digite apenas o número, sem texto nem unidade."
    dv.showErrorMessage = True
    ws.add_data_validation(dv)
    dv.add(rng)


def cf_equal(ws, rng, pairs):
    for text, color in pairs:
        ws.conditional_formatting.add(rng, CellIsRule(operator="equal", formula=['"%s"' % text], fill=PatternFill("solid", bgColor=color, fgColor=color)))


def cf_ok(ws, rng):
    ok = PatternFill("solid", bgColor=GREEN, fgColor=GREEN)
    warn = PatternFill("solid", bgColor=YELLOW, fgColor=YELLOW)
    ws.conditional_formatting.add(rng, CellIsRule(operator="equal", formula=['"OK"'], fill=ok))
    ws.conditional_formatting.add(rng, CellIsRule(operator="notEqual", formula=['"OK"'], fill=warn))


def cf_status(ws, rng, first):
    """Verde para OK, amarelo para qualquer outro texto, nada para a célula vazia.
    first é a primeira célula do intervalo, com a referência relativa ou absoluta que a regra pede."""
    cf_equal(ws, rng, [("OK", GREEN)])
    ws.conditional_formatting.add(rng, FormulaRule(formula=[f'AND(LEN({first})>0,{first}<>"OK")'],
                                                   fill=PatternFill("solid", bgColor=YELLOW, fgColor=YELLOW)))


def note(ws, ref, text):
    c = Comment(text, "Modelo")
    c.width, c.height = 260, 90
    ws[ref].comment = c


def summary(ws, row, text, formula, lab_to, val, fmt=None, val_merge=None):
    label(ws, f"B{row}", text, merge=f"B{row}:{lab_to}{row}", h="right")
    calc(ws, f"{val}{row}", formula, fmt=fmt, merge=val_merge)
    ws.row_dimensions[row].height = 19.5


def instrucoes(wb, titulo, sub, secoes):
    """Preenche a primeira aba (Instruções). secoes = [(título da seção, [(rótulo, texto), ...]), ...].
    Use os rótulos "Amarelo-claro" e "Cinza" para a legenda de cores: a linha recebe a cor."""
    ws = wb.active
    ws.title = "Instruções"
    widths(ws, {"A": 2, "B": 24, "C": 96, "D": 2})
    title(ws, titulo, sub, "C")
    r = 4
    for sec, linhas in secoes:
        band(ws, r, sec, "C", sz=11)
        r += 1
        for k, v in linhas:
            vbg = INPUT if k == "Amarelo-claro" else GRAY if str(k).startswith("Cinza") else None
            put(ws, f"B{r}", k, f=font(10, True), bg=GRAY, h="center" if isinstance(k, int) or len(str(k)) <= 3 else "left")
            put(ws, f"C{r}", v, bg=vbg)
            ws.row_dimensions[r].height = 19.5 if len(str(v)) <= 100 else 31.5
            r += 1
        r += 1
    setup(ws, INK, f"B1:C{r}", landscape=False)
    return ws


def checklist(ws, titulo, sub, itens, tab=None, primeira=5):
    """Aba de checklist: item, situação (Sim / Parcial / Não), observação, e o percentual atendido.
    itens = [texto, ...] ou [(grupo, texto), ...]."""
    widths(ws, {"A": 2, "B": 6, "C": 22, "D": 70, "E": 14, "F": 40, "G": 2})
    title(ws, titulo, sub, "F")
    r = primeira - 1
    for col, t in zip("BCDEF", ["#", "Grupo", "Item", "Situação", "Observação"]):
        head(ws, f"{col}{r}", t)
    ws.row_dimensions[r].height = 21.75
    for n, it in enumerate(itens, 1):
        r += 1
        grupo, texto = it if isinstance(it, tuple) else ("", it)
        num(ws, f"B{r}", n)
        put(ws, f"C{r}", grupo, bg=GRAY)
        put(ws, f"D{r}", texto)
        inp(ws, f"E{r}", h="center")
        inp(ws, f"F{r}")
        ws.row_dimensions[r].height = 19.5 if len(texto) <= 75 else 31.5
    ini, fim = primeira, r
    dv_list(ws, f"E{ini}:E{fim}", ["Sim", "Parcial", "Não"], "Sim, Parcial ou Não.")
    cf_equal(ws, f"E{ini}:E{fim}", [("Sim", GREEN), ("Parcial", YELLOW), ("Não", RED)])
    r += 2
    summary(ws, r, "Itens respondidos", f'=COUNTA(E{ini}:E{fim})', "D", "E")
    summary(ws, r + 1, "Percentual atendido (Parcial vale metade)",
            f'=IF(COUNTA(E{ini}:E{fim})=0,"",(COUNTIF(E{ini}:E{fim},"Sim")+COUNTIF(E{ini}:E{fim},"Parcial")/2)/COUNTA(E{ini}:E{fim}))',
            "D", "E", fmt="0%")
    summary(ws, r + 2, "Situação",
            f'=IF(COUNTA(E{ini}:E{fim})<{len(itens)},"Incompleto",IF(COUNTIF(E{ini}:E{fim},"Não")=0,"OK","Há itens a resolver"))',
            "D", "E")
    cf_ok(ws, f"E{r + 2}")
    ws.freeze_panes = f"B{primeira}"
    setup(ws, tab or HUE["teal"][0], f"B1:F{r + 2}")
    return ws

