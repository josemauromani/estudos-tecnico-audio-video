# -*- coding: utf-8 -*-
"""Gera o painel (index.html) a partir de index_dados.py, de index.css e dos estudos que já existem.

    python build_index.py <saida/index.html>

Os números de cada cartão (módulos, minutos, figuras, abas) são contados nos arquivos do estudo.
Um estudo que ainda não tem corpo aparece como "em preparação", sem botões.
"""
import os
import re
import sys
from html import escape

AQUI = os.path.dirname(os.path.abspath(__file__))
RAIZ = os.path.dirname(AQUI)
sys.path.insert(0, AQUI)
from index_dados import FASES, GRUPOS, TRILHAS  # noqa: E402

OUT = sys.argv[1]

ARQ = {}
for linha in open(os.path.join(AQUI, "estudos.txt"), encoding="utf-8"):
    if linha.strip() and not linha.startswith("#"):
        nome, pasta, html, xlsx = linha.strip().split("|")
        ARQ[nome] = (pasta, html, xlsx)

ESTUDOS = [e for _, _, lista in GRUPOS for e in lista]
POR_NOME = {e["nome"]: e for e in ESTUDOS}
assert set(POR_NOME) == set(ARQ), set(POR_NOME) ^ set(ARQ)


def fatos(nome):
    """(módulos, minutos, figuras, abas) do estudo, ou None se ele ainda não foi gerado."""
    pasta, html, xlsx = ARQ[nome]
    corpo = os.path.join(AQUI, nome + "_body.html")
    if not (os.path.exists(corpo) and os.path.exists(os.path.join(RAIZ, pasta, html))):
        return None
    t = open(corpo, encoding="utf-8").read()
    mod = len(re.findall(r'<p class="eyebrow">Módulo', t))
    m = re.search(r"Leitura de cerca de (\d+) minutos", t)
    minutos = int(m.group(1)) if m else 0
    fig = len(re.findall(r"<figure\b", t))
    abas = 0
    caminho = os.path.join(RAIZ, pasta, xlsx)
    if os.path.exists(caminho):
        from openpyxl import load_workbook
        abas = len(load_workbook(caminho, read_only=True).sheetnames)
    return mod, minutos, fig, abas


FATOS = {n: fatos(n) for n in ARQ}
prontos = [n for n in ARQ if FATOS[n]]
n_mod = sum(FATOS[n][0] for n in prontos)
n_min = sum(FATOS[n][1] for n in prontos)
n_xlsx = sum(1 for n in prontos if FATOS[n][3])
N = len(ESTUDOS)
POR_EXTENSO = {21: "Vinte e um"}


def link_html(nome):
    pasta, html, _ = ARQ[nome]
    return f"{pasta}/{html}"


def tiles(e):
    return "".join(f'<span class="{c}">{escape(t)}</span>' for c, t in e["tiles"])


def cartao(e):
    nome = e["nome"]
    pasta, html, xlsx = ARQ[nome]
    f = FATOS[nome]
    tag = '\n      <span class="tag">Com vídeo</span>' if e.get("video") else ""
    if f:
        planilha = f" · planilha com {f[3]} abas" if f[3] else ""
        fatos_txt = f"{f[0]} módulos · {f[1]} min · {f[2]} figuras{planilha}"
        botoes = f'<a class="btn" href="{pasta}/{html}">Abrir o treinamento</a>'
        if f[3]:
            botoes += f'\n        <a class="btn ghost" href="{pasta}/{xlsx}" download>Planilha .xlsx</a>'
        acoes = f'<div class="actions">\n        {botoes}\n      </div>'
        cls = "card"
    else:
        fatos_txt = "em preparação"
        acoes = ""
        cls = "card soon"
    return f"""    <article class="{cls}" id="{nome}">
      <div class="tiles" aria-hidden="true">{tiles(e)}</div>{tag}
      <h3>{escape(e["titulo"])}</h3>
      <p>{escape(e["desc"])}</p>
      <p class="ask"><b>Responde a:</b> {escape(e["ask"])}</p>
      <p class="facts">{fatos_txt}</p>
      {acoes}
      <label class="done-box"><input type="checkbox" data-k="{nome}"><span>Estudo concluído</span></label>
    </article>"""


# ------------------------------------------------------------------ figura da sequência
def caixa(nome, x, y, w, rotulo, sub, ink=False, h=72):
    e = POR_NOME[nome]
    cls, tx = ("bx-ink", " t-ground") if ink else ("bx", "")
    mu = " t-ground" if ink else " mu"
    q = "".join(f'<rect class="f-{c.split()[0]}" x="{x + 12 + 20 * i}" y="{y + 10}" width="18" height="18"/>'
                for i, (c, _) in enumerate(e["tiles"]))
    return (f'<a href="#{nome}"><rect class="{cls}" x="{x}" y="{y}" width="{w}" height="{h}"/>{q}'
            f'<text class="b{tx}" x="{x + 12}" y="{y + 46}" font-size="12.5">{escape(rotulo)}</text>'
            f'<text class="{mu.strip()}" x="{x + 12}" y="{y + 62}" font-size="10.5">{escape(sub)}</text></a>')


def faixa(nome, y, rotulo, sub):
    e = POR_NOME[nome]
    q = "".join(f'<rect class="f-{c.split()[0]}" x="{24 + 20 * i}" y="{y + 16}" width="18" height="18"/>'
                for i, (c, _) in enumerate(e["tiles"]))
    x0 = 24 + 20 * len(e["tiles"]) + 10
    return (f'<a href="#{nome}"><rect class="bx-ink" x="10" y="{y}" width="880" height="50"/>{q}'
            f'<text class="b t-ground" x="{x0}" y="{y + 30}" font-size="13.5">{escape(rotulo)}</text>'
            f'<text class="t-ground" x="876" y="{y + 30}" font-size="11.5" text-anchor="end">{escape(sub)}</text></a>')


def seta(d):
    return f'<path class="ln" d="{d}" marker-end="url(#ar)"/>'


X4 = [20, 240, 460, 680]
svg = []
svg.append(faixa("som", 12, "Som, frequência e decibel", "a base: o vocabulário de todos os estudos de áudio"))
svg.append('<text class="mono mu" x="20" y="92" font-size="10">ÁUDIO · O CAMINHO DO SINAL, DA FONTE AO PÚBLICO</text>')
for x, (n, r, s) in zip(X4, [("mic", "Microfones", "captar a fonte"), ("cabos", "Cabos e conexões", "levar o sinal sem ruído"),
                             ("mesa", "Mesa de som", "rotear e misturar"), ("caixas", "Caixas e PA", "entregar ao público")]):
    svg.append(caixa(n, x, 102, 200, r, s))
for x in X4[:3]:
    svg.append(seta(f"M{x + 201} 138 H{x + 217}"))
svg.append(seta("M560 176 V200"))
svg.append('<rect class="bx-out" x="10" y="204" width="880" height="118"/>')
svg.append('<text class="mono mu" x="876" y="222" font-size="10" text-anchor="end">NA MESA, CANAL POR CANAL</text>')
for x, (n, r, s) in zip(X4, [("ganho", "Ganho e fader", "sinal limpo, com folga"), ("eq", "Equalização", "cada som no seu espaço"),
                             ("din", "Dinâmica e efeitos", "controlar e dar ambiência"), ("ret", "Monitores e retornos", "o que o músico ouve")]):
    svg.append(caixa(n, x, 232, 200, r, s))
for x in X4[:3]:
    svg.append(seta(f"M{x + 201} 268 H{x + 217}"))
for x, w, (n, r, s) in zip([20, 315], [270, 270], [("fb", "Microfonia", "quebrar o ciclo caixa → microfone"),
                                                   ("semfio", "Microfones sem fio", "frequências, antenas e pilhas")]):
    svg.append(caixa(n, x, 346, w, r, s))
svg.append(caixa("pass", 610, 346, 270, "Passagem de som e mixagem", "junta tudo no dia do evento", ink=True))
svg.append(seta("M780 324 V342"))
svg.append('<text x="770" y="338" font-size="10.5" text-anchor="end" class="mu">canais prontos</text>')

svg.append('<text class="mono mu" x="20" y="452" font-size="10">VÍDEO · O CAMINHO DA IMAGEM, DA LUZ AO ESPECTADOR</text>')
for x, (n, r, s) in zip(X4, [("luz", "Iluminação", "luz para a câmera ver"), ("expo", "Exposição", "íris, obturador e ganho"),
                             ("lente", "Lentes, foco e cor", "nitidez e cor certa"), ("quadro", "Enquadramento", "o que mostrar, e como")]):
    svg.append(caixa(n, x, 462, 200, r, s))
for x in X4[:3]:
    svg.append(seta(f"M{x + 201} 498 H{x + 217}"))
svg.append(seta("M780 536 V553 H120 V560"))
svg.append('<text x="450" y="548" font-size="10.5" text-anchor="middle" class="mu">a imagem de cada câmera</text>')
for x, (n, r, s) in zip(X4, [("sinal", "Sinais e cabos de vídeo", "todos no mesmo formato"), ("corte", "Corte ao vivo", "escolher a imagem"),
                             ("live", "Transmissão ao vivo", "imagem e som na internet")]):
    svg.append(caixa(n, x, 564, 200, r, s, ink=(n == "live")))
for x in X4[:2]:
    svg.append(seta(f"M{x + 201} 600 H{x + 217}"))
svg.append(caixa("proj", 680, 564, 200, "Projeção, telão e mídia", "letras e vídeos na sala"))
svg.append('<path class="ln dash" d="M678 600 H664" stroke-dasharray="4 3"/>')
svg.append('<path class="ln" d="M880 382 H890 V660 H560 V640" marker-end="url(#ar)" stroke-dasharray="4 3"/>')
svg.append('<text x="720" y="654" font-size="10.5" text-anchor="middle" class="mu">mixagem de áudio própria para a transmissão</text>')
svg.append(faixa("oper", 676, "Operação do evento", "roteiro, equipe, energia e imprevistos: organiza todos os outros"))
SVG_ALT = ("Mapa dos estudos. No topo, Som, frequência e decibel é a base de todos os estudos de áudio. "
           "No caminho do áudio: Microfones captam a fonte, Cabos e conexões levam o sinal, a Mesa de som roteia e mistura, "
           "e Caixas e PA entregam o som ao público. Dentro da mesa, canal por canal: Ganho e fader, Equalização, "
           "Dinâmica e efeitos, e Monitores e retornos. Microfonia e Microfones sem fio tratam dois problemas do sistema. "
           "Passagem de som e mixagem junta tudo no dia do evento. No caminho do vídeo: Iluminação, Exposição, "
           "Lentes, foco e cor, e Enquadramento formam a imagem de cada câmera. Essa imagem segue por Sinais e cabos de vídeo "
           "para o Corte ao vivo e para a Transmissão ao vivo, que também recebe uma mixagem de áudio própria. "
           "Projeção, telão e mídia mostra letras e vídeos na sala. Na base, Operação do evento organiza todos os outros.")
figura = (f'<svg viewBox="0 0 900 740" role="img" aria-label="{SVG_ALT}">\n<defs><marker id="ar" viewBox="0 0 10 10" refX="9" refY="5" '
          'markerWidth="7" markerHeight="7" orient="auto"><path class="ah" d="M0 0 L10 5 L0 10 z"/></marker></defs>\n'
          + "\n".join(svg) + "\n</svg>")

# ------------------------------------------------------------------ blocos
needs = "\n".join(f'    <li><a href="#{e["nome"]}"><span>{escape(e["need"])}</span><b>{escape(e["titulo"])}</b></a></li>' for e in ESTUDOS)
grupos = "\n\n".join(
    f'  <h4>{escape(t)}</h4>\n  <div class="cards{" few" if cls else ""}">\n' + "\n\n".join(cartao(e) for e in lista) + "\n  </div>"
    for t, cls, lista in GRUPOS)
seq = "\n".join(
    f'        <tr><td class="num">{i}</td><td><a href="#{e["nome"]}"><strong>{escape(e["titulo"])}</strong></a></td>'
    f'<td>{escape(e["recebe"])}</td><td>{escape(e["entrega"])}</td></tr>' for i, e in enumerate(ESTUDOS, 1))
trilhas = "\n".join(
    f'    <div class="trail"><h3>{escape(t)}</h3><p>{escape(d)}</p><ol>\n'
    + "\n".join(f'      <li><a href="#{n}">{escape(POR_NOME[n]["titulo"])}</a></li>' for n in nomes)
    + "\n    </ol></div>" for t, d, nomes in TRILHAS)
fases_head = "".join(f'<th class="ph">{escape(t)}</th>' for _, t in FASES)
fases = "\n".join(
    f'        <tr><td><a href="#{e["nome"]}"><strong>{escape(e["titulo"])}</strong></a></td>'
    + "".join(f'<td class="ph on">●</td>' if k in e["fases"] else '<td class="ph"></td>' for k, _ in FASES)
    + "</tr>" for e in ESTUDOS)

VIDEOS = [
    ("ganho", "Ganho não é volume", "AQOI3QxrfV1vMrEH_dlkVhoNcAx8x_GgAk8yhTEXuuWLt2KZTyoD0Tq6I6cTxUJYMhCHPa5UzgNjgo.mp4",
     "Ganho é a entrada da caixa d'água, fader é a torneira. 1 min 15 s."),
    ("fb", "Microfonia: por que a caixa apita?", "AQMMfS8_JQpGs0iAgD1D5V_J1fcfQwkT8GHxzVZ1OwARrg52_BHof53GRgw5aovTFLWMkXUv.mp4",
     "O ciclo entre o microfone e a caixa, como a câmera que filma a própria TV. 1 min."),
]
videos = "\n".join(
    f'    <div class="vid"><video controls preload="metadata" src="docs/tutoriais/audio/{arq}"></video>'
    f'<div><h3>{escape(t)}</h3><p>{escape(d)}</p><p>Estudo: <a href="#{n}">{escape(POR_NOME[n]["titulo"])}</a></p></div></div>'
    for n, t, arq, d in VIDEOS)

horas = max(1, round(n_min / 60))
resumo = (f"<li>{N} estudos</li>\n    <li>{n_mod} módulos</li>\n    <li>Cerca de {horas} horas de leitura</li>\n    <li>{n_xlsx} planilhas de apoio</li>"
          if len(prontos) == N else
          f"<li>{N} estudos no mapa</li>\n    <li>{len(prontos)} prontos</li>\n    <li>{n_mod} módulos</li>\n    <li>{n_xlsx} planilhas de apoio</li>")
css = open(os.path.join(AQUI, "index.css"), encoding="utf-8").read()
FONTS = ("https://fonts.googleapis.com/css2?family=Barlow+Semi+Condensed:wght@500;600;700"
         "&family=IBM+Plex+Mono:wght@400;500&family=IBM+Plex+Sans:ital,wght@0,400;0,500;0,600;1,400&display=swap")
ext = POR_EXTENSO.get(N, str(N))

html = f"""<!DOCTYPE html>
<html lang="pt-BR">
<head>
<meta charset="utf-8">
<meta name="viewport" content="width=device-width, initial-scale=1, viewport-fit=cover">
<title>Técnico em áudio e vídeo</title>
<link rel="stylesheet" href="{FONTS}">
<style>
{css}</style>
</head>
<body>
<div class="page" lang="pt-BR">

<header class="mast">
  <p class="kicker">Estudos · Técnico em áudio e vídeo</p>
  <h1>Som e imagem ao vivo na prática</h1>
  <p class="lead">{ext} estudos autoinstrucionais para operar mesa de som, câmeras, corte e transmissão em igrejas e eventos ao vivo. Cada um tem um treinamento para ler e uma planilha para usar no evento. Escolha pelo que você precisa fazer, siga o caminho do sinal ou a trilha da sua função.</p>
  <ul class="meta">
    {resumo}
  </ul>
</header>

<main>

<!-- ============================ ESCOLHA ============================ -->
<section id="escolha">
  <p class="eyebrow">Por onde começar</p>
  <h2>O que você precisa fazer?</h2>
  <p>Cada estudo responde a uma pergunta diferente. Escolha a necessidade, e o painel leva ao estudo correspondente.</p>
  <ul class="needs">
{needs}
  </ul>
</section>

<!-- ============================ ESTUDOS ============================ -->
<section id="estudos">
  <p class="eyebrow">Os estudos</p>
  <h2>{ext} estudos, quatro grupos</h2>
  <p>Marque cada estudo quando terminar. As marcações ficam salvas neste navegador.</p>
  <div class="bar"><span><b id="done-count">0</b> de {N} estudos concluídos</span><span class="track" aria-hidden="true"><i id="done-bar"></i></span><button id="done-reset" type="button">Limpar marcações</button></div>

{grupos}
</section>

<!-- ============================ TRILHAS ============================ -->
<section id="trilhas">
  <p class="eyebrow">Trilhas</p>
  <h2>A ordem de leitura para cada função</h2>
  <p>Ninguém precisa ler tudo de uma vez. Cada trilha põe em ordem os estudos de que uma função precisa primeiro. Os estudos são independentes: quem já domina um assunto pode pular.</p>
  <div class="trails">
{trilhas}
  </div>
</section>

<!-- ============================ SEQUÊNCIA ============================ -->
<section id="sequencia">
  <p class="eyebrow">Visão de conjunto</p>
  <h2>Como os estudos se encadeiam</h2>
  <p>Os estudos seguem o caminho do sinal. No áudio, o som sai da fonte, é captado pelo microfone, viaja pelo cabo, é tratado e misturado na mesa e chega ao público pelas caixas. No vídeo, a luz ilumina a cena, a câmera forma a imagem, o sinal viaja até a mesa de corte e de lá segue para o telão e para a transmissão. O estudo de som, frequência e decibel dá o vocabulário de todos os estudos de áudio. A passagem de som junta os estudos de áudio no dia do evento, e a operação do evento organiza a equipe, o roteiro e os imprevistos de todas as áreas.</p>

  <figure>
    <div class="scroll">
{figura}
    </div>
    <figcaption><b>Mapa dos estudos.</b> Cada caixa leva ao cartão do estudo. As setas seguem o sinal: o que um estudo entrega é o que o seguinte recebe.</figcaption>
  </figure>

  <div class="tbl">
    <table>
      <thead><tr><th style="width:6%">#</th><th style="width:24%">Estudo</th><th style="width:32%">O que recebe</th><th>O que entrega</th></tr></thead>
      <tbody>
{seq}
      </tbody>
    </table>
  </div>
</section>

<!-- ============================ FASES ============================ -->
<section id="fases">
  <p class="eyebrow">No dia do evento</p>
  <h2>Quando cada estudo é usado</h2>
  <p>Um evento tem cinco momentos: o planejamento, a montagem, a passagem, o evento e a revisão depois dele. A tabela mostra em que momento o conteúdo de cada estudo entra em uso. Ela serve para revisar, na véspera, só o que será preciso em cada etapa.</p>
  <div class="tbl">
    <table>
      <thead><tr><th>Estudo</th>{fases_head}</tr></thead>
      <tbody>
{fases}
      </tbody>
    </table>
  </div>
</section>

<!-- ============================ VÍDEOS ============================ -->
<section id="videos">
  <p class="eyebrow">Tutoriais em vídeo</p>
  <h2>Os vídeos da série</h2>
  <p>Os vídeos curtos da pasta <code>docs/tutoriais/audio</code> apresentam a ideia central de dois estudos. Assista antes de ler: o treinamento retoma a mesma analogia e vai além. A transcrição está em <a href="docs/tutoriais/audio/transcricoes.md">transcricoes.md</a>.</p>
  <div class="vids">
{videos}
  </div>
  <div class="note">
    <h4>Pontos de atenção</h4>
    <p>Os valores dos estudos são faixas típicas e pontos de partida. Cada equipamento tem os seus limites: confira sempre o manual. Em tudo o que envolve eletricidade, peso suspenso, trabalho em altura e volume alto, a conduta segura vem antes do resultado.</p>
  </div>
</section>

</main>

<footer>Material de estudo. Os exemplos, os cenários e os valores dos treinamentos são ilustrativos e foram criados para fins didáticos. Eles não substituem o manual dos equipamentos nem as normas de segurança aplicáveis.</footer>
</div>

<script>
(function () {{
  var KEY = 'av-estudos-concluidos';
  function read() {{ try {{ return JSON.parse(localStorage.getItem(KEY) || '{{}}') || {{}}; }} catch (e) {{ return {{}}; }} }}
  function write(v) {{ try {{ localStorage.setItem(KEY, JSON.stringify(v)); }} catch (e) {{}} }}

  var boxes = Array.prototype.slice.call(document.querySelectorAll('.done-box input'));
  var count = document.getElementById('done-count');
  var bar = document.getElementById('done-bar');
  var state = read();

  function refresh() {{
    var n = 0;
    boxes.forEach(function (b) {{
      var card = b.closest('.card');
      if (b.checked) {{ n++; card.classList.add('done'); }} else {{ card.classList.remove('done'); }}
      state[b.getAttribute('data-k')] = b.checked;
    }});
    count.textContent = n;
    bar.style.width = (n / boxes.length * 100) + '%';
  }}
  boxes.forEach(function (b) {{
    if (state[b.getAttribute('data-k')]) {{ b.checked = true; }}
    b.addEventListener('change', function () {{ refresh(); write(state); }});
  }});
  refresh();
  document.getElementById('done-reset').addEventListener('click', function () {{
    boxes.forEach(function (b) {{ b.checked = false; }});
    refresh(); write(state);
  }});
}})();
</script>
</body>
</html>
"""
open(OUT, "w", encoding="utf-8").write(html)
print("ok", OUT, f"{len(prontos)}/{N} estudos prontos")
